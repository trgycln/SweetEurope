import os
import glob
import re
import fitz  # PyMuPDF
import docx
import olefile
import struct
import json
import sys

sys.stdout.reconfigure(encoding='utf-8')

def extract_doc_raw(filepath):
    if filepath.endswith('.docx'):
        try:
            doc = docx.Document(filepath)
            txt = '\n'.join([p.text for p in doc.paragraphs])
            for t in doc.tables:
                for r in t.rows:
                    txt += '\n' + ' | '.join([c.text.strip() for c in r.cells])
            return txt
        except Exception as e:
            return f"Error docx: {e}"
    try:
        ole = olefile.OleFileIO(filepath)
        word_stream = ole.openstream('WordDocument').read()
        flags = struct.unpack_from('<H', word_stream, 0x000A)[0]
        table_stream = ole.openstream('1Table' if (flags & 0x0200) else '0Table').read()
        fcClx = struct.unpack_from('<I', word_stream, 0x01A2)[0]
        lcbClx = struct.unpack_from('<I', word_stream, 0x01A6)[0]
        clx = table_stream[fcClx : fcClx + lcbClx]
        pos = 0
        while pos < len(clx):
            clxt = clx[pos]
            pos += 1
            if clxt == 1:
                pos += 2 + struct.unpack_from('<H', clx, pos)[0]
            elif clxt == 2:
                cb = struct.unpack_from('<I', clx, pos)[0]
                pos += 4
                plc = clx[pos : pos + cb]
                n = (cb - 4) // 12
                cps = [struct.unpack_from('<I', plc, i * 4)[0] for i in range(n + 1)]
                pcds_offset = (n + 1) * 4
                full = []
                for i in range(n):
                    fc_val = struct.unpack_from('<I', plc, pcds_offset + i * 8 + 2)[0]
                    cnt = cps[i+1] - cps[i]
                    fc = fc_val & 0x3FFFFFFF
                    if not (fc_val & 0x40000000):
                        full.append(word_stream[fc : fc + cnt * 2].decode('utf-16le', errors='ignore'))
                    else:
                        full.append(word_stream[fc//2 : fc//2 + cnt].decode('cp1252', errors='ignore'))
                ole.close()
                return ''.join(full)
        ole.close()
    except Exception as e:
        return f"Error ole: {e}"
    return ''

def extract_pdf_raw(filepath):
    try:
        doc = fitz.open(filepath)
        txt = ""
        for page in doc:
            txt += page.get_text("text") + "\n"
        doc.close()
        return txt
    except Exception as e:
        return f"Error pdf: {e}"

def analyze():
    spec_files = glob.glob("dokuments/FO Ürün Spektleri/**/*.doc*", recursive=True)
    label_files = glob.glob("dokuments/Ürün Etiketleri/**/*.pdf", recursive=True)

    print(f"Total Spec files found: {len(spec_files)}")
    print(f"Total Label PDF files found: {len(label_files)}")

    results = {
        "specs": {
            "total_files": len(spec_files),
            "files_with_temp": 0,
            "files_with_after_opening": 0,
            "unique_temperatures": {},
            "storage_sentences": {},
            "after_opening_matches": []
        },
        "labels": {
            "total_files": len(label_files),
            "files_with_temp": 0,
            "files_with_after_opening": 0,
            "unique_temperatures": {},
            "storage_sentences": {},
            "after_opening_matches": []
        }
    }

    # 1. Analyze SPECS
    for sf in spec_files:
        text = extract_doc_raw(sf)
        fname = os.path.basename(sf)
        
        # Temp matches like 18 - 22 °C or 20°C
        temp_matches = re.findall(r'(\d+\s*[-–]\s*\d+\s*°?\s*C|\d+\s*°\s*C)', text, re.IGNORECASE)
        # Storage section
        storage_m = re.findall(r'(?:storage conditions|storage|saklama)[^\n\r|\x07]{0,200}', text, re.IGNORECASE)
        # Shelf life
        shelf_m = re.findall(r'(?:shelf life)[^\n\r|\x07]{0,100}', text, re.IGNORECASE)
        # After opening
        after_m = re.findall(r'[^\n\r|\x07]{0,60}(?:opening|opened|açıldık|nach öffnen|geöffnet|consume within|gün içinde)[^\n\r|\x07]{0,60}', text, re.IGNORECASE)

        if temp_matches:
            results["specs"]["files_with_temp"] += 1
            for tm in temp_matches:
                t_clean = re.sub(r'\s+', ' ', tm).strip()
                results["specs"]["unique_temperatures"][t_clean] = results["specs"]["unique_temperatures"].get(t_clean, 0) + 1

        for sm in storage_m:
            s_clean = re.sub(r'\s+', ' ', sm).strip()
            if len(s_clean) > 10:
                results["specs"]["storage_sentences"][s_clean] = results["specs"]["storage_sentences"].get(s_clean, 0) + 1

        if after_m:
            results["specs"]["files_with_after_opening"] += 1
            for am in after_m:
                results["specs"]["after_opening_matches"].append({
                    "file": fname,
                    "match": am.strip()
                })

    # 2. Analyze LABELS
    for lf in label_files:
        text = extract_pdf_raw(lf)
        fname = os.path.basename(lf)

        temp_matches = re.findall(r'(\d+\s*[-–]\s*\d+\s*°?\s*C|\d+\s*°\s*C)', text, re.IGNORECASE)
        storage_m = re.findall(r'(?:saklayınız|muhafaza|aufbewahren|store|lagerung)[^\n\r]{0,150}', text, re.IGNORECASE)
        after_m = re.findall(r'[^\n\r]{0,60}(?:opening|opened|açıldık|nach öffnen|geöffnet|consume within|gün içinde|tüketiniz|buzdolabı|kühlschrank)[^\n\r]{0,60}', text, re.IGNORECASE)

        if temp_matches:
            results["labels"]["files_with_temp"] += 1
            for tm in temp_matches:
                t_clean = re.sub(r'\s+', ' ', tm).strip()
                results["labels"]["unique_temperatures"][t_clean] = results["labels"]["unique_temperatures"].get(t_clean, 0) + 1

        for sm in storage_m:
            s_clean = re.sub(r'\s+', ' ', sm).strip()
            if len(s_clean) > 10:
                results["labels"]["storage_sentences"][s_clean] = results["labels"]["storage_sentences"].get(s_clean, 0) + 1

        if after_m:
            results["labels"]["files_with_after_opening"] += 1
            for am in after_m:
                results["labels"]["after_opening_matches"].append({
                    "file": fname,
                    "match": am.strip()
                })

    with open("scratch/temp_shelf_analysis.json", "w", encoding="utf-8") as f:
        json.dump(results, f, ensure_ascii=False, indent=2)

    print("SUCCESS: Results written to scratch/temp_shelf_analysis.json")
    print(f"Specs with Temp: {results['specs']['files_with_temp']} / {results['specs']['total_files']}")
    print(f"Specs Unique Temp values: {results['specs']['unique_temperatures']}")
    print(f"Specs with After-Opening: {results['specs']['files_with_after_opening']}")
    print("-----------------------------------------")
    print(f"Labels with Temp: {results['labels']['files_with_temp']} / {results['labels']['total_files']}")
    print(f"Labels Unique Temp values: {results['labels']['unique_temperatures']}")
    print(f"Labels with After-Opening: {results['labels']['files_with_after_opening']}")

if __name__ == '__main__':
    analyze()
